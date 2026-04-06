#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""Word 文档拆分实现。"""

import json
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Tuple

from ..core.logger import Logger
from .contracts import ContextSegment, DocumentInfo, ImageInfo, StructuredElement
from .doc_profile import classify_word_document, normalize_elements_by_doc_type
from .splitter_base import (
    DocumentSplitter,
    generate_smart_image_name_simple,
    generate_smart_image_name_with_llm,
)


class WordSplitter(DocumentSplitter):
    """Word文档拆分器 - 使用语义分割"""
    
    def __init__(self, chunk_size: int = 1000, chunk_overlap: int = 200,
                 use_llm_naming: bool = False, config: dict = None):
        """
        初始化Word拆分器
        
        Args:
            chunk_size: 每个chunk的字符数（默认1000，适合中文文档）
            chunk_overlap: chunk重叠字符数（默认200，保证上下文连贯）
            use_llm_naming: 是否使用 LLM 智能生成图片名称（默认False）
            config: 处理器配置字典（None 则自动加载）
        """
        super().__init__(chunk_size, chunk_overlap, use_llm_naming, config=config)
        try:
            from docx import Document
            from docx.oxml.table import CT_Tbl
            from docx.oxml.text.paragraph import CT_P
            from docx.table import _Cell, Table
            from docx.text.paragraph import Paragraph
            self.Document = Document
            self.CT_Tbl = CT_Tbl
            self.CT_P = CT_P
            self.Paragraph = Paragraph
        except ImportError as e:
            raise ImportError(f"缺少依赖 python-docx: {e}。请安装: pip install python-docx") from e
    
    def process(self, word_path: Path, output_dir: Path) -> DocumentInfo:
        """
        处理Word文档
        
        Args:
            word_path: Word文件路径
            output_dir: 输出目录（直接在此目录下创建 chunks 和 images）
            
        Returns:
            文档信息
        """
        Logger.info(f"处理Word文档: {word_path.name}")
        
        # 创建输出目录结构
        doc_dir = output_dir
        chunks_dir = doc_dir / "chunks"
        images_dir = doc_dir / "images"
        
        chunks_dir.mkdir(parents=True, exist_ok=True)
        images_dir.mkdir(parents=True, exist_ok=True)
        
        # 加载文档（验证文件有效性）
        try:
            doc = self.Document(word_path)
        except Exception as e:
            error_msg = f"无法打开Word文档（可能已损坏或不是有效的 .docx 文件）: {word_path.name}"
            Logger.error(error_msg)
            Logger.error(f"错误详情: {e}", indent=1)
            raise ValueError(error_msg) from e
        
        # 提取文本和图片信息（带占位符 + 结构化元素）
        full_text, image_infos, structured_elements = self._extract_text_with_image_placeholders(doc, word_path)
        doc_profile = classify_word_document(structured_elements)
        structured_elements = normalize_elements_by_doc_type(structured_elements, doc_profile)
        Logger.info(
            f"文档类型识别: {doc_profile.doc_type} | "
            f"段落={doc_profile.paragraph_count} 表格={doc_profile.table_count} "
            f"图片={doc_profile.image_count} 重复提示={doc_profile.prompt_repeat_count}",
            indent=1,
        )
        
        # 提取并保存图片（使用智能命名）
        images_count = self._save_word_images_with_smart_names(word_path, images_dir, image_infos, doc_dir)
        
        # 更新文本中的占位符，使用智能文件名和相对路径格式
        full_text = self._update_image_placeholders(full_text, image_infos)
        structured_elements = self._replace_element_image_placeholders(structured_elements, image_infos)

        # 拆分为chunks
        chunks = self._build_structured_word_chunks(word_path.name, structured_elements, doc_profile=doc_profile)

        # 保存chunks
        for chunk in chunks:
            chunk_file = chunks_dir / f"chunk_{chunk.index:04d}.txt"
            with open(chunk_file, 'w', encoding='utf-8') as f:
                f.write(chunk.content)
        
        # 保存chunks索引
        self._save_chunks_index(chunks, doc_dir)
        
        # 保存文档信息
        doc_info = DocumentInfo(
            filename=word_path.name,
            format=f"Word/{doc_profile.doc_type}",
            total_pages=len(doc.sections),
            total_chunks=len(chunks),
            total_images=images_count,
            created_at=datetime.now().isoformat(),
            output_dir=str(doc_dir),
            doc_type=doc_profile.doc_type,
            profile=doc_profile.to_dict(),
        )
        
        self._save_doc_info(doc_info, doc_dir)
        
        Logger.success(f"完成: {len(chunks)} 个chunks, {images_count} 张图片")
        return doc_info
    
    def _extract_text_with_image_placeholders(self, doc, word_path: Path) -> Tuple[str, List[ImageInfo], List[StructuredElement]]:
        """
        提取Word文档文本，并在图片位置插入占位符
        使用两次遍历策略获取更完整的上下文
        增强：识别标题样式
        
        Returns:
            (完整文本, 图片信息列表, 结构化元素列表)
        """
        import zipfile
        
        # 获取所有图片文件名
        image_files = []
        with zipfile.ZipFile(word_path) as docx_zip:
            for file_info in docx_zip.filelist:
                if file_info.filename.startswith('word/media/'):
                    image_files.append(Path(file_info.filename).name)
        
        # 第一步：收集所有段落文本、图片位置和格式信息
        all_elements: List[StructuredElement] = []
        table_index = 0
        
        for element in doc.element.body:
            if isinstance(element, self.CT_P):
                # 检查是否有图片
                element_xml = str(element.xml)
                has_image = ('a:blip' in element_xml or 'blip' in element_xml.lower() or 
                            'v:imagedata' in element_xml or 'imagedata' in element_xml.lower())
                
                paragraph = self.Paragraph(element, doc)
                para_text = paragraph.text.strip()
                
                # 去除段落内的明显重复（如"GVG报名GVG报名GVG报名"）
                para_text = self._remove_internal_duplicates(para_text)
                
                # 识别标题/列表样式
                is_heading, heading_level, is_list_item, list_level = self._extract_word_paragraph_structure(
                    paragraph, para_text
                )
                
                all_elements.append(
                    StructuredElement(
                        element_type='paragraph',
                        text=para_text,
                        has_image=has_image,
                        is_heading=is_heading,
                        heading_level=heading_level,
                        is_list_item=is_list_item,
                        list_level=list_level,
                    )
                )
                
            elif isinstance(element, self.CT_Tbl):
                # 处理表格
                table_payload = self._extract_table_payload(element)
                if table_payload["text"]:
                    table_index += 1
                    all_elements.append(
                        StructuredElement(
                            element_type='table',
                            text=table_payload["text"],
                            table_headers=table_payload["headers"],
                            table_rows=table_payload["rows"],
                            table_index=table_index,
                            caption=table_payload["caption"],
                        )
                    )
        
        # 从配置读取上下文提取参数
        context_config = self.config.get('doc_splitter', {}).get('context_extraction', {})
        
        lookback = context_config.get('word_lookback_paragraphs', 8)
        lookahead = context_config.get('word_lookahead_paragraphs', 5)
        max_length = context_config.get('max_length', 600)
        
        # 第二步：处理图片，提取更丰富的上下文
        image_infos = []
        image_index = 0
        para_list = []
        current_position = 0
        
        for idx, element in enumerate(all_elements):
            elem_type = element.element_type
            content = element.text
            has_image = element.has_image
            is_heading = element.is_heading
            heading_level = element.heading_level
            if has_image and image_index < len(image_files):
                # 提取上下文：可配置的段落数
                context_before_parts = []
                context_after_parts = []
                context_segments_before = []  # 新增：带格式的上下文片段
                context_segments_after = []   # 新增：带格式的上下文片段
                
                # 向前看最多 lookback 个有内容的元素
                distance = 0
                for i in range(idx - 1, max(0, idx - lookback) - 1, -1):
                    if all_elements[i].text:  # 如果有内容
                        elem_text = all_elements[i].text
                        elem_is_heading = all_elements[i].is_heading
                        elem_heading_level = all_elements[i].heading_level
                        
                        context_before_parts.append(elem_text)
                        
                        # 创建上下文片段
                        segment = ContextSegment(
                            text=elem_text,
                            distance=distance,
                            is_heading=elem_is_heading,
                            heading_level=elem_heading_level
                        )
                        context_segments_before.append(segment)
                        distance += 1
                
                # 向后看最多 lookahead 个有内容的元素
                distance = 0
                # 如果当前段落有文本，也加入 context_after
                if content:
                    context_after_parts.insert(0, content)
                    segment = ContextSegment(
                        text=content,
                        distance=0,
                        is_heading=is_heading,
                        heading_level=heading_level
                    )
                    context_segments_after.append(segment)
                    distance = 1
                
                for i in range(idx + 1, min(len(all_elements), idx + lookahead + 1)):
                    if all_elements[i].text:
                        elem_text = all_elements[i].text
                        elem_is_heading = all_elements[i].is_heading
                        elem_heading_level = all_elements[i].heading_level
                        
                        context_after_parts.append(elem_text)
                        
                        # 创建上下文片段
                        segment = ContextSegment(
                            text=elem_text,
                            distance=distance,
                            is_heading=elem_is_heading,
                            heading_level=elem_heading_level
                        )
                        context_segments_after.append(segment)
                        distance += 1
                
                # 合并上下文（兼容旧版）
                context_before = ' '.join(context_before_parts)
                context_after = ' '.join(context_after_parts)
                
                # 限制上下文长度（可配置）
                context_before = context_before[-max_length:] if len(context_before) > max_length else context_before
                context_after = context_after[:max_length] if len(context_after) > max_length else context_after
                
                # 创建图片信息
                original_filename = f"img_{image_index + 1:03d}_{image_files[image_index]}"
                image_info = ImageInfo(
                    page_number=0,
                    image_index=image_index + 1,
                    position=current_position,
                    context_before=context_before,
                    context_after=context_after,
                    original_filename=original_filename,
                    smart_filename="",
                    context_segments_before=context_segments_before,  # 新增
                    context_segments_after=context_segments_after     # 新增
                )
                image_infos.append(image_info)
                
                # 构建文本
                if content:
                    para_list.append(content)
                    current_position += len(content) + 2
                
                placeholder = f"[图片: images/{original_filename}]"
                para_list.append(placeholder)
                current_position += len(placeholder) + 2
                
                image_index += 1
                
            elif content:
                # 普通段落或表格
                para_list.append(content)
                current_position += len(content) + 2
        
        full_text = "\n\n".join(para_list)
        structured_elements: List[StructuredElement] = []
        image_cursor = 0
        for element in all_elements:
            structured_elements.append(element)
            if element.has_image and image_cursor < len(image_infos):
                image_info = image_infos[image_cursor]
                image_cursor += 1
                structured_elements.append(
                    StructuredElement(
                        element_type='paragraph',
                        text=f"[图片: images/{image_info.original_filename}]",
                    )
                )
        return full_text, image_infos, structured_elements
    
    def _remove_internal_duplicates(self, text: str) -> str:
        """
        去除文本内部的明显重复
        例如：'GVG报名GVG报名GVG报名' -> 'GVG报名'
        
        Args:
            text: 原始文本
            
        Returns:
            去重后的文本
        """
        return self._deduplicate_text(text)
    
    def _extract_table_payload(self, table_element) -> Dict[str, Any]:
        """提取表格文本及结构化信息，尽量保留合并单元格和多级表头。"""
        ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        rows: List[List[str]] = []

        for row in table_element.findall('./w:tr', namespaces=ns):
            cells: List[str] = []
            for cell in row.findall('./w:tc', namespaces=ns):
                cell_text = self._compact_whitespace(''.join(cell.itertext()))

                grid_span = 1
                try:
                    grid_span_nodes = cell.findall('./w:tcPr/w:gridSpan', namespaces=ns)
                    if grid_span_nodes:
                        grid_span = max(1, int(grid_span_nodes[0].get(f'{{{ns["w"]}}}val', '1')))
                except Exception:
                    grid_span = 1

                repeated_value = cell_text
                if not repeated_value:
                    repeated_value = ""
                cells.extend([repeated_value] * grid_span)

            if any(cell.strip() for cell in cells):
                rows.append(cells)

        rows = self._normalize_table_rows(rows)
        header_depth = self._infer_table_header_depth(rows)
        header_rows = rows[:header_depth] if header_depth else []
        headers = self._merge_table_headers(header_rows)
        data_rows = rows[header_depth:] if len(rows) > header_depth else []
        if not data_rows and rows:
            data_rows = rows[1:] if len(rows) > 1 else rows
        text_rows = [' | '.join(cell for cell in row if cell or len(row) == 1) for row in rows]
        caption = " | ".join(headers[:3]).strip() if headers else ""
        return {
            "text": '\n'.join(text_rows) if text_rows else "",
            "headers": headers,
            "rows": data_rows,
            "caption": caption,
            "header_depth": header_depth,
        }
    
    def _save_word_images_with_smart_names(self, word_path: Path, images_dir: Path,
                                            image_infos: List[ImageInfo], doc_dir: Path) -> int:
        """
        提取并保存Word文档图片，使用智能命名
        
        Returns:
            图片数量
        """
        import zipfile
        
        image_count = 0
        
        with zipfile.ZipFile(word_path) as docx_zip:
            media_files = [f for f in docx_zip.filelist if f.filename.startswith('word/media/')]
            
            for idx, file_info in enumerate(media_files):
                if idx < len(image_infos):
                    image_info = image_infos[idx]
                    
                    # 提取图片数据
                    image_data = docx_zip.read(file_info.filename)
                    image_filename = Path(file_info.filename).name
                    image_ext = Path(image_filename).suffix
                    
                    if self.use_llm_naming and self.llm:
                        result = generate_smart_image_name_with_llm(
                            image_info.context_before,
                            image_info.context_after,
                            max_length=self.image_max_length,
                            llm=self.llm,
                            context_segments_before=image_info.context_segments_before,
                            context_segments_after=image_info.context_segments_after,
                            image_data=image_data,
                            original_filename=image_filename,
                        )
                        smart_name, description = result if isinstance(result, tuple) else (result, "")
                        image_info.description = description
                    else:
                        smart_name = generate_smart_image_name_simple(
                            image_info.context_before,
                            image_info.context_after,
                            max_length=self.image_max_length
                        )
                    
                    image_info.smart_filename = f"{smart_name}_{idx + 1:03d}{image_ext}"
                    
                    # 保存图片（使用智能命名）
                    image_path = images_dir / image_info.smart_filename
                    with open(image_path, 'wb') as f:
                        f.write(image_data)
                    
                    Logger.info(f"保存图片: {image_info.smart_filename}", indent=1)
                    image_count += 1
        
        # 保存图片信息索引
        self._save_images_index(image_infos, doc_dir)
        
        return image_count
    
    def _save_images_index(self, image_infos: List[ImageInfo], doc_dir: Path):
        """保存图片信息索引"""
        index_file = doc_dir / "images_index.json"
        index_data = {
            "total_images": len(image_infos),
            "images": [
                {
                    "page_number": info.page_number if info.page_number > 0 else "N/A",
                    "image_index": info.image_index,
                    "smart_filename": info.smart_filename,
                    "original_filename": info.original_filename,
                    "description": getattr(info, 'description', ''),
                    "context_before": info.context_before[:100],
                    "context_after": info.context_after[:100],
                }
                for info in image_infos
            ]
        }
        
        with open(index_file, 'w', encoding='utf-8') as f:
            json.dump(index_data, f, ensure_ascii=False, indent=2)




